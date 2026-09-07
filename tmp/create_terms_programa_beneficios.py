from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("Documentos/Terminos_Programa_de_Beneficios_Futura_Tecno.docx")

DARK = "16181D"
LIME = "C8E048"
DEEP_LIME = "5D6B14"
MUTED = "5D6673"
LIGHT = "F3F5F7"


def set_font(run, name="Arial", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths_dxa):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "180")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_layout = OxmlElement("w:tblLayout")
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_pr.append(tbl_layout)
    grid = table._tbl.tblGrid
    for gc, width in zip(grid.gridCol_lst, widths_dxa):
        gc.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def add_text(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, size=10.8, color=DARK, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r, size=10.8, color=DARK)
    else:
        r = p.add_run(text)
        set_font(r, size=10.8, color=DARK)
    return p


def add_list(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    set_font(r, size=10.8, color=DARK)
    return p


def add_h1(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(13)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, size=15, color=DEEP_LIME, bold=True)
    return p


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT)
    set_cell_margins(cell, top=130, start=180, bottom=130, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    set_font(r, size=10.8, color=DEEP_LIME, bold=True)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(body)
    set_font(r, size=10.5, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.8)
    normal.font.color.rgb = RGBColor.from_string(DARK)

    for name in ("Heading 1", "Heading 2", "Heading 3"):
        st = doc.styles[name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        st.font.color.rgb = RGBColor.from_string(DEEP_LIME)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("FUTURATECNO  |  PROGRAMA DE BENEFICIOS")
    set_font(r, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Futura Tecno · Términos del Programa de Beneficios · Versión 1.0")
    set_font(r, size=8.5, color=MUTED)

    # First-page title block using standard business brief + restrained brand override.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("FUTURATECNO")
    set_font(r, size=15, color=LIME, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("TÉRMINOS DEL PROGRAMA DE BENEFICIOS")
    set_font(r, size=22, color=DARK, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(20)
    r = p.add_run("Borrador para publicación · Versión 1.0 · Vigencia a definir")
    set_font(r, size=10.5, color=MUTED, italic=True)

    add_callout(
        doc,
        "Resumen del programa",
        "Los usuarios registrados acumulan 1 punto por cada USD 10 de productos efectivamente pagados. Cada punto puede canjearse por un descuento equivalente a USD 0,20, convertido a pesos al momento de confirmar la compra.",
    )

    add_h1(doc, "1. Objeto y aceptación")
    add_text(doc, "El Programa de Beneficios de Futura Tecno (el “Programa”) permite a los usuarios registrados acumular y canjear puntos conforme a estos términos. Al participar, realizar compras o canjear puntos, la persona usuaria acepta estas condiciones y las políticas aplicables del sitio futuratecno.com.ar.")

    add_h1(doc, "2. Personas participantes")
    add_text(doc, "Podrán participar las personas con una cuenta activa y datos de registro válidos en futuratecno.com.ar. La cuenta, los puntos y los beneficios son personales, intransferibles y no pueden venderse, cederse ni convertirse en dinero.")

    add_h1(doc, "3. Obtención de puntos")
    add_text(doc, "Los puntos se acreditarán exclusivamente cuando el pago de una compra se encuentre aprobado y confirmado por el medio de pago integrado. La relación de acumulación será la siguiente:")
    add_list(doc, "1 punto por cada USD 10 de valor neto de productos pagados.")
    add_list(doc, "El cálculo se realizará sobre el importe de los productos, sin incluir envío, cargos financieros, impuestos específicos, propinas ni importes reembolsados.")
    add_list(doc, "Se computarán únicamente puntos enteros; las fracciones inferiores a USD 10 no se acumulan ni se trasladan a una compra posterior.")
    add_list(doc, "Los puntos se mostrarán inicialmente como pendientes y se acreditarán cuando el pago figure como aprobado.")

    add_h1(doc, "4. Ajustes, anulaciones y reintegros")
    add_text(doc, "Futura Tecno podrá revertir, descontar o anular puntos cuando una compra se cancele, se rechace, se desconozca, se reintegre total o parcialmente, presente un contracargo, o se detecte un error técnico o uso indebido. La reversión se aplicará en proporción al importe reembolsado cuando corresponda.")

    add_h1(doc, "5. Consulta y vigencia de los puntos")
    add_text(doc, "La persona usuaria podrá consultar su saldo y movimientos desde la sección “Mis puntos” de su cuenta. Cada punto vencerá a los doce (12) meses contados desde la fecha de su acreditación. Futura Tecno podrá informar campañas con una vigencia distinta; esa condición especial se mostrará de forma expresa antes de participar.")

    add_h1(doc, "6. Canje de puntos")
    add_text(doc, "Los puntos podrán utilizarse como descuento al finalizar una compra elegible. El canje estará sujeto a estas reglas:")
    add_list(doc, "1 punto equivale a un descuento de USD 0,20.")
    add_list(doc, "El valor del descuento se convertirá a pesos argentinos según la cotización oficial de venta informada en el checkout al confirmar el pedido.")
    add_list(doc, "El canje mínimo es de 20 puntos por pedido.")
    add_list(doc, "El descuento por puntos no podrá superar el 20% del valor de los productos del pedido ni cubrir el costo de envío, cargos financieros o impuestos aplicables.")
    add_list(doc, "Salvo que una promoción indique expresamente lo contrario, los puntos no son acumulables con otros cupones o descuentos promocionales.")
    add_list(doc, "Los puntos se debitarán al confirmar el pedido. Si el pago no se aprueba o el pedido se cancela antes de su preparación, se reintegrarán los puntos utilizados.")

    add_h1(doc, "7. Uso responsable")
    add_text(doc, "Queda prohibido crear cuentas duplicadas, manipular precios, pagos o promociones, automatizar operaciones, usar datos de terceros sin autorización o realizar cualquier conducta fraudulenta. Frente a indicios razonables de fraude o incumplimiento, Futura Tecno podrá suspender la cuenta, retener puntos pendientes y anular puntos obtenidos indebidamente, sin perjuicio de otras medidas aplicables.")

    add_h1(doc, "8. Cambios, suspensión o finalización")
    add_text(doc, "Futura Tecno podrá modificar, suspender o finalizar el Programa por razones operativas, comerciales, técnicas, legales o de seguridad. Cuando el cambio afecte sustancialmente los puntos ya acreditados, se comunicará por medios razonables con antelación. En caso de finalización, se informará un plazo razonable para utilizar los puntos vigentes, salvo que una causa de seguridad, fraude o exigencia legal requiera una medida inmediata.")

    add_h1(doc, "9. Datos personales")
    add_text(doc, "Los datos necesarios para administrar el Programa se tratarán conforme a las políticas de privacidad de Futura Tecno y la normativa aplicable. Los movimientos de puntos podrán asociarse a compras aprobadas, pedidos, cancelaciones y devoluciones para operar el Programa y prevenir usos indebidos.")

    add_h1(doc, "10. Contacto y normativa aplicable")
    add_text(doc, "Las consultas sobre puntos, canjes o movimientos podrán realizarse por los canales oficiales publicados en futuratecno.com.ar. Estos términos se regirán por las leyes de la República Argentina, sin perjuicio de los derechos que correspondan a las personas consumidoras.")

    add_callout(
        doc,
        "Nota de publicación",
        "Antes de publicar, completar la fecha de entrada en vigencia y revisar este texto con asesoramiento legal y contable propio, especialmente la mecánica de conversión a pesos, promociones y tratamiento impositivo.",
    )
    doc.save(OUT)


if __name__ == "__main__":
    build()
